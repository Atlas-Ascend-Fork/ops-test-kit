#!/usr/bin/env python3
# -*- coding: UTF-8 -*-
# Copyright (c) 2026 Huawei Technologies Co., Ltd.
# This program is free software, you can redistribute it and/or modify it under the terms and conditions of
# CANN Open Software License Agreement Version 2.0 (the "License").
# Please refer to the License for details. You may not use this file except in compliance with the License.
# THIS SOFTWARE IS PROVIDED ON AN "AS IS" BASIS, WITHOUT WARRANTIES OF ANY KIND, EITHER EXPRESS OR IMPLIED,
# INCLUDING BUT NOT LIMITED TO NON-INFRINGEMENT, MERCHANTABILITY, OR FITNESS FOR A PARTICULAR PURPOSE.
# See LICENSE in the root of the software repository for the full text of the License.
"""openpyxl 工作簿解析。"""

__all__ = ["read_xlsx_rows", "resolved_sheet"]


def read_xlsx_rows(path, sheet=None) -> tuple:
    """Read an .xlsx workbook into (header, rows) of stripped strings.

    sheet: worksheet name; default first worksheet. Cells are coerced to
    str (None -> "") and stripped. Fully-empty rows are dropped, matching
    the CSV reader's empty-line filtering. data_only=True reads cached
    formula values; numeric cells become their decimal string form, so
    shape/dtype columns should be formatted as Text in the workbook.
    """
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover
        raise ImportError("openpyxl is required to read .xlsx inputs; install it with `pip install openpyxl`") from exc

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet is not None:
            if sheet not in wb.sheetnames:
                raise ValueError(f"Sheet '{sheet}' not found in {path}; available: {wb.sheetnames}")
            ws = wb[sheet]
        else:
            ws = wb[wb.sheetnames[0]]
        rows = []
        for raw in ws.iter_rows(values_only=True):
            row = [("" if cell is None else str(cell)).strip() for cell in raw]
            if any(cell != "" for cell in row):
                rows.append(row)
        if not rows:
            sheet_name = sheet or wb.sheetnames[0]
            raise ValueError(f"Empty sheet '{sheet_name}' in {path}")
        return rows[0], rows[1:]
    finally:
        wb.close()


def resolved_sheet(path, sheet=None):
    """实际读取的工作表名（xlsx 未指定 --sheet 时为首个表名；csv/zip 为 None）。"""
    if not path.lower().endswith((".xlsx", ".xlsm")):
        return None
    try:
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            return sheet if sheet is not None else wb.sheetnames[0]
        finally:
            wb.close()
    except Exception:
        return None
